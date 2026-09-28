from __future__ import annotations

from io import BytesIO

import pandas as pd

from analysis import score_distribution, score_text, summary_metrics, unit_summary


def build_briefing_docx(segments: list[tuple[str, pd.DataFrame]]) -> bytes:
    """Create an editable Word briefing from the currently selected Eastern data."""
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    page = document.sections[0]
    page.page_width = Inches(8.5)
    page.page_height = Inches(11)
    page.top_margin = Inches(0.72)
    page.bottom_margin = Inches(0.68)
    page.left_margin = Inches(0.72)
    page.right_margin = Inches(0.72)

    def set_font(style_name: str, size: int, bold: bool = False) -> None:
        style = document.styles[style_name]
        style.font.name = "Noto Sans SC"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.insert(0, fonts)
        fonts.set(qn("w:eastAsia"), "Noto Sans SC")

    set_font("Normal", 10)
    set_font("Title", 20, True)
    set_font("Heading 1", 15, True)
    set_font("Heading 2", 11, True)
    document.styles["Normal"].paragraph_format.space_after = Pt(5)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(10)

    def shade(cell) -> None:
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F2F5FA")
        properties.append(shading)

    def add_chart(cell, title: str, values: list[tuple[str, float]], scale: float, count: bool = False) -> None:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        shade(cell)
        heading = cell.paragraphs[0]
        heading.style = "Heading 2"
        heading.add_run(title)
        for label, value in values:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(label))
            run.font.size = Pt(9)

            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            bar_units = max(0, min(18, round(float(value) / scale * 18))) if scale else 0
            bar = p.add_run("━" * max(1, bar_units))
            bar.font.color.rgb = RGBColor(65, 105, 201)
            bar.font.size = Pt(13)
            number = p.add_run(f"  {int(value) if count else score_text(value)}")
            number.font.size = Pt(9)
        cell.add_paragraph().paragraph_format.space_after = Pt(1)

    document.add_paragraph("东航型别教员技能评估简报", style="Title")

    for index, (title, frame) in enumerate(segments):
        if index:
            document.add_page_break()
        document.add_paragraph(title, style="Heading 1")
        metrics = summary_metrics(frame, pd.DataFrame())
        lead = document.add_paragraph()
        lead.paragraph_format.space_after = Pt(12)
        lead.add_run(f"参加{metrics['评估人数']}人，平均分数{score_text(metrics['平均模拟机得分'])}。")

        distribution = score_distribution(frame)
        distribution_values = [
            (str(row["分数区间"]), float(row["人数"]))
            for _, row in distribution.iterrows()
            if int(row["人数"]) > 0
        ]
        units = unit_summary(frame)
        unit_values = [(str(row["所属单位"]), float(row["平均分"])) for _, row in units.iterrows()]

        panels = document.add_table(rows=1, cols=2)
        panels.autofit = False
        for cell in panels.rows[0].cells:
            cell.width = Inches(3.48)
        add_chart(
            panels.cell(0, 0), "得分分布", distribution_values,
            max((value for _, value in distribution_values), default=1), count=True,
        )
        add_chart(panels.cell(0, 1), "单位平均分", unit_values, 100)
        document.add_paragraph().paragraph_format.space_after = Pt(0)

    output = BytesIO()
    document.save(output)
    return output.getvalue()
